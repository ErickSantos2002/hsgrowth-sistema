"""
Gera `criterios.py` a partir da planilha da consultoria.

Digitar 26 critérios com três rubricas cada, à mão, é convite a erro de peso —
e peso errado desloca o score de todas as reuniões sem aparecer na tela.

Uso (a planilha fica fora do container, então copie antes):
    docker cp "Documentação/Avaliacao_Calls_HealthSafety_alinhada_a_matriz.xlsm" \
        hsgrowth-api-local:/tmp/avaliacao.xlsm
    docker exec -e PYTHONPATH=/app -w /app hsgrowth-api-local \
        python scripts/extrair_criterios.py /tmp/avaliacao.xlsm
"""
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

PRIMEIRA_LINHA = 2
ULTIMA_LINHA = 27
PRIMEIRA_FAIXA = 36
ULTIMA_FAIXA = 39

DESTINO = "app/services/avaliacao_reuniao/criterios.py"


def texto(valor) -> str:
    return str(valor).strip().replace("\n", " ") if valor is not None else ""


def cabecalho_do_arquivo() -> str:
    return '''"""
A régua da consultoria — 26 critérios, peso e rubrica de três níveis.

GERADO por `scripts/extrair_criterios.py` a partir de
`Documentação/Avaliacao_Calls_HealthSafety_alinhada_a_matriz.xlsm`.
Não editar à mão: mudou a régua, rode o script de novo e suba a VERSAO.

A versão fica gravada em cada avaliação. Sem isso, mudar um peso reescreveria
o passado: uma reunião avaliada em setembro apareceria com outra nota em
novembro, e ninguém saberia por quê.
"""
from dataclasses import dataclass


@dataclass(frozen=True)
class Criterio:
    id: str
    bloco: str
    titulo: str
    peso: float
    rubrica: tuple  # textos das notas 0, 1 e 2


@dataclass(frozen=True)
class Faixa:
    minimo: float
    veredito: str


VERSAO = "2026-09"

BLOCOS = ("Abertura", "Diagnóstico", "Demonstração", "Fechamento")

# Abaixo disto a reunião não recebe score comparável: transcrição curta ou
# parcial puxa a nota para baixo por falta de conversa, não por falta de
# técnica. Na planilha da consultoria, uma call de 9 minutos ficou assim.
COBERTURA_MINIMA = 0.70

VEREDITO_PARCIAL = "Call parcial — não comparar"


'''


def rodape_do_arquivo() -> str:
    return '''def criterio_por_id(criterio_id: str):
    """Devolve o critério, ou None se o id não existir na régua."""
    return next((c for c in CRITERIOS if c.id == criterio_id), None)
'''


def main(caminho: str) -> None:
    ws = openpyxl.load_workbook(caminho, data_only=True)["Criterios"]

    criterios = []
    for linha in ws.iter_rows(min_row=PRIMEIRA_LINHA, max_row=ULTIMA_LINHA, values_only=True):
        if not linha[0]:
            continue
        criterios.append({
            "id": texto(linha[0]),
            "bloco": texto(linha[1]),
            "titulo": texto(linha[2]),
            "peso": float(linha[4]),
            "rubrica": (texto(linha[5]), texto(linha[6]), texto(linha[7])),
        })

    faixas = []
    for linha in ws.iter_rows(min_row=PRIMEIRA_FAIXA, max_row=ULTIMA_FAIXA, values_only=True):
        if linha[1] is None:
            continue
        faixas.append({"minimo": float(linha[1]), "veredito": texto(linha[3])})
    faixas.sort(key=lambda f: f["minimo"], reverse=True)

    soma = sum(c["peso"] for c in criterios)
    if len(criterios) != 26 or round(soma, 2) != 100:
        raise SystemExit(f"Planilha inesperada: {len(criterios)} criterios, pesos somam {soma}")

    with open(DESTINO, "w", encoding="utf-8") as f:
        f.write(cabecalho_do_arquivo())
        f.write("CRITERIOS: tuple = (\n")
        for c in criterios:
            f.write("    Criterio(\n")
            f.write(f'        id={c["id"]!r},\n')
            f.write(f'        bloco={c["bloco"]!r},\n')
            f.write(f'        titulo={c["titulo"]!r},\n')
            f.write(f'        peso={c["peso"]!r},\n')
            f.write("        rubrica=(\n")
            for nivel in c["rubrica"]:
                f.write(f"            {nivel!r},\n")
            f.write("        ),\n")
            f.write("    ),\n")
        f.write(")\n\n")
        f.write("FAIXAS: tuple = (\n")
        for faixa in faixas:
            f.write(f'    Faixa(minimo={faixa["minimo"]!r}, veredito={faixa["veredito"]!r}),\n')
        f.write(")\n\n\n")
        f.write(rodape_do_arquivo())

    print(f"{DESTINO}: {len(criterios)} criterios, pesos somam {soma}, {len(faixas)} faixas")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "/tmp/avaliacao.xlsm")
